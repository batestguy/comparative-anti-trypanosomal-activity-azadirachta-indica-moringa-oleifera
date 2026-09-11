#!/usr/bin/env python3
"""Combine report.docx + objectives-methods.docx into one send-ready document.

Appends the methods companion as an appendix (page break + retitled heading),
carrying over any missing table styles (e.g. flextable's "Normal Table").

Usage:
  python3 analysis/R/combine_report.py
Output:
  analysis/reports/full-report.docx
"""
import copy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[2]
REPORT = REPO / "analysis" / "reports" / "report.docx"
METHODS = REPO / "analysis" / "reports" / "objectives-methods.docx"
OUT = REPO / "analysis" / "reports" / "full-report.docx"


def ensure_style(dst_doc, src_doc, style_name):
    """Copy a style definition from src to dst if dst lacks it."""
    try:
        dst_doc.styles[style_name]
        return False
    except KeyError:
        pass
    for s in src_doc.styles:
        if s.name == style_name:
            dst_doc.styles.element.append(copy.deepcopy(s.element))
            return True
    return False


def main():
    report = Document(str(REPORT))
    methods = Document(str(METHODS))

    # flextable tables need their table style present
    for st in ("Normal Table", "Table Grid"):
        if ensure_style(report, methods, st):
            print(f"carried over style: {st}")

    # page break, then appendix heading (retitled first Heading 1)
    br = report.add_paragraph()
    run = br.add_run()
    brk = OxmlElement("w:br")
    brk.set(qn("w:type"), "page")
    run._r.append(brk)

    first_h1_done = False
    for el in list(methods.element.body):
        if el.tag != qn("w:p") and el.tag != qn("w:tbl"):
            continue
        new_el = copy.deepcopy(el)
        if not first_h1_done and el.tag == qn("w:p"):
            p = None
            # find style of source paragraph
            pPr = el.find(qn("w:pPr"))
            style = pPr.find(qn("w:pStyle")).get(qn("w:val")) if pPr is not None and pPr.find(qn("w:pStyle")) is not None else ""
            try:
                sname = methods.styles.element.xpath(
                    f'.//w:style[@w:styleId="{style}"]/w:name/@w:val')[0]
            except IndexError:
                sname = ""
            if sname == "heading 1":
                for r in new_el.findall(qn("w:r")):
                    for t in r.findall(qn("w:t")):
                        if t.text and t.text.strip():
                            t.text = "Appendix: " + t.text.lstrip()
                            first_h1_done = True
                            break
                    if first_h1_done:
                        break
        report.element.body.append(new_el)

    # refresh the "Appendix" label font size consistency is inherited; just save
    report.save(str(OUT))
    print("wrote", OUT,
          "| paragraphs:", len(report.paragraphs),
          "| tables:", len(report.tables))


if __name__ == "__main__":
    main()
